import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = pptx.Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_BG = RGBColor(11, 17, 32)           # #0B1120 Deep Dark Navy
    C_CARD = RGBColor(19, 29, 49)         # #131D31 Card background
    C_CARD_BORDER = RGBColor(30, 58, 95)  # #1E3A5F Border
    C_CARD_ALT = RGBColor(15, 23, 42)     # #0F172A Darker card
    C_CYAN = RGBColor(56, 189, 248)       # #38BDF8 Electric Cyan / Accent
    C_GREEN = RGBColor(16, 185, 129)      # #10B981 Emerald Green / Success
    C_ORANGE = RGBColor(245, 158, 11)     # #F59E0B Warning / Stat
    C_PURPLE = RGBColor(168, 85, 247)     # #A855F7 AI / Innovation
    C_RED = RGBColor(239, 68, 68)         # #EF4444 Alert
    C_WHITE = RGBColor(248, 250, 252)     # #F8FAFC White primary text
    C_MUTED = RGBColor(148, 163, 184)     # #94A3B8 Slate muted text
    C_LIGHT_BG = RGBColor(23, 37, 84)     # #172554 Subtle banner

    def set_slide_background(slide, color=C_BG):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title, category="RIDEHACK'26 | CYPHER-X", subtext=""):
        # Category / Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_CYAN
        p_cat.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.55))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_WHITE
        p_title.font.name = "Arial"

        # Subtext if provided
        if subtext:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.22), Inches(11.7), Inches(0.35))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtext
            p_sub.font.size = Pt(12)
            p_sub.font.color.rgb = C_MUTED
            p_sub.font.name = "Arial"

    def add_card(slide, left, top, width, height, bg_color=C_CARD, border_color=C_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # -------------------------------------------------------------
    # SLIDE 1: COVER SLIDE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Top Event Tag
    tag_card = add_card(slide1, Inches(0.8), Inches(0.7), Inches(4.5), Inches(0.4), bg_color=RGBColor(15, 23, 42), border_color=C_CYAN)
    tf_tag = tag_card.text_frame
    tf_tag.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_tag = tf_tag.paragraphs[0]
    p_tag.alignment = PP_ALIGN.CENTER
    p_tag.text = "⚡ RIDEHACK'26 — IIC JUIT STARTUP CHALLENGE"
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = C_CYAN
    p_tag.font.name = "Arial"

    # Big Title
    t_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.5), Inches(1.2))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "SATYAMEVA AEGIS"
    p_t.font.size = Pt(44)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE
    p_t.font.name = "Arial"

    # Subtitle
    sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(8.5), Inches(0.6))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Decentralized Trust Protocol & Academic Credential Verification Platform"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = C_CYAN
    p_sub.font.name = "Arial"

    # Elevator Pitch Card
    pitch_card = add_card(slide1, Inches(0.8), Inches(3.2), Inches(8.0), Inches(1.4), bg_color=C_CARD, border_color=C_CARD_BORDER)
    tf_p = pitch_card.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = Inches(0.3)
    tf_p.margin_top = Inches(0.2)
    p_p = tf_p.paragraphs[0]
    p_p.text = "Eliminating the ₹4,000 Cr academic document counterfeit black market. Transforming university credentials into tamper-proof, self-sovereign digital assets verified in under 2 seconds without accounts."
    p_p.font.size = Pt(13)
    p_p.font.color.rgb = C_WHITE
    p_p.font.name = "Arial"

    # Logo/Graphic on right if available
    icon_path = os.path.abspath("static/img/icon-512.png")
    if os.path.exists(icon_path):
        slide1.shapes.add_picture(icon_path, Inches(9.5), Inches(1.5), width=Inches(3.0))

    # Feature Highlights Grid at bottom
    badges = [
        ("⚡ 2-Sec Instant Verify", "Lifetime QR scan with institutional crest"),
        ("⛓️ Custom PoW Ledger", "Zero gas fees + PBKDF2 dynamic salting"),
        ("🤖 Gemini 2.0 Flash AI", "Automated marksheet parsing & bulk CSV"),
        ("🏛️ JUIT / JIIT Pilot Ready", "Multi-campus isolation & 100% E2E passed")
    ]
    for i, (b_title, b_sub) in enumerate(badges):
        bx = Inches(0.8 + i * 2.95)
        bc = add_card(slide1, bx, Inches(4.8), Inches(2.8), Inches(1.1), bg_color=RGBColor(15, 23, 42), border_color=RGBColor(30, 58, 95))
        tf_b = bc.text_frame
        tf_b.margin_left = Inches(0.15)
        tf_b.margin_top = Inches(0.15)
        p1 = tf_b.paragraphs[0]
        p1.text = b_title
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = C_CYAN
        p2 = tf_b.add_paragraph()
        p2.text = b_sub
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = C_MUTED

    # Footer Metadata
    ft_box = slide1.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.6))
    tf_ft = ft_box.text_frame
    p_ft = tf_ft.paragraphs[0]
    p_ft.text = "TEAM: CYPHER-X (Team ID: 10)   |   FOUNDER: Praneeth Sai   |   HOST: Jaypee University of Information Technology (JUIT Waknaghat)   |   STAGE: Working Prototype"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = C_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: THE PROBLEM & MARKET CRISIS
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "THE CRISIS IN CREDENTIAL INTEGRITY & VERIFICATION",
               category="Problem Statement & Market Pain",
               subtext="How manual processes, paper fraud, and fragmented verification cost billions annually.")

    problems = [
        ("🚨 40,000+ Fake Degrees / Year",
         "The Fake Degree Epidemic",
         "A booming counterfeit certificate black market in India exceeds ₹4,000 Crores. Forged engineering, medical, and management degrees regularly slip undetected into the corporate workforce.",
         C_RED),
        ("⏳ 2 to 4 Weeks Turnaround",
         "Agonizing Verification Delays",
         "Corporate recruiters and Background Verification (BGV) agencies wait weeks for manual postal replies and phone calls to university registrars, freezing corporate onboarding pipelines.",
         C_ORANGE),
        ("💸 ₹500 to ₹2,000 Per Applicant",
         "Crushing Operational Costs",
         "Employers and background screening companies spend thousands per candidate on manual registrar checks, while universities waste massive administrative hours answering enquiries.",
         C_PURPLE),
        ("🔓 Zero Cryptographic Tamper Resistance",
         "Paper & Static PDF Vulnerability",
         "Traditional paper certificates and static PDFs offer zero cryptographic protection. Furthermore, centralized university databases remain vulnerable to rogue internal DB tampering.",
         C_CYAN)
    ]

    for i, (p_tag, p_head, p_desc, p_col) in enumerate(problems):
        col = i % 2
        row = i // 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.7 + row * 2.2)
        card = add_card(slide2, x, y, Inches(5.8), Inches(2.0), bg_color=C_CARD, border_color=C_CARD_BORDER)
        tf = card.text_frame
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.25)

        p1 = tf.paragraphs[0]
        p1.text = p_tag
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = p_col

        p2 = tf.add_paragraph()
        p2.text = p_head
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        p3 = tf.add_paragraph()
        p3.text = p_desc
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = C_MUTED

    # Bottom Stat Callout Banner
    banner = add_card(slide2, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.7), bg_color=RGBColor(23, 37, 84), border_color=C_CYAN)
    tf_b = banner.text_frame
    tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    p_b.text = "🎯 NATIONAL REALITY: Over 40 Million Higher Education students in India are stuck with fragile paper certificates and zero self-sovereign digital ownership."
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = C_CYAN

    # -------------------------------------------------------------
    # SLIDE 3: THE SOLUTION - SATYAMEVA AEGIS
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "OUR SOLUTION: INSTANT, DECENTRALIZED TRUST",
               category="Value Proposition & Core Capabilities",
               subtext="Transforming degrees into cryptographically anchored, self-sovereign digital assets.")

    solutions = [
        ("🏛️ Permanent Institutional QR",
         "Campus-Branded Verification Anchor",
         "High-contrast QR code embedded with the official university crest (JUIT Waknaghat / JIIT Noida) using Error Correction Level H (~30%). Lifetime-valid, permanent, and surrounded by a 360° cyber radar HUD ring.",
         C_CYAN),
        ("🛡️ Zero Silent Passes (Anti-Tamper)",
         "Dynamic PBKDF2 Re-Derivation",
         "Our verification engine doesn't just read database records; it re-derives the cryptographic hash from scratch on every request and audits the blockchain. Modifying even 1 SQL letter trips an instant TAMPERED alert.",
         C_GREEN),
        ("🤖 Google Gemini 2.0 Flash AI",
         "Intelligent Registrar Automation",
         "Eliminates hours of manual data entry. Registrars upload certificates; Gemini AI parses name, roll number, course, and CGPA in 2 seconds. Includes atomic all-or-nothing validation for bulk CSV batches.",
         C_PURPLE)
    ]

    for i, (s_tag, s_head, s_desc, s_col) in enumerate(solutions):
        x = Inches(0.8 + i * 3.95)
        card = add_card(slide3, x, Inches(1.7), Inches(3.8), Inches(3.3), bg_color=C_CARD, border_color=C_CARD_BORDER)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)

        p1 = tf.paragraphs[0]
        p1.text = s_tag
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = s_col

        p2 = tf.add_paragraph()
        p2.text = s_head
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        p3 = tf.add_paragraph()
        p3.text = s_desc
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = C_MUTED

    # Bottom 3 Metric Badges
    metrics = [
        ("< 2 SECONDS", "Instant Verification Speed", "Reduced from 2 to 4 weeks (99.9% faster)", C_CYAN),
        ("> 98% SAVINGS", "Corporate BGV Cost Slashed", "From ₹1,500 down to ₹15 per check", C_GREEN),
        ("100% CERTAINTY", "Cryptographic Tamper Detection", "Zero silent passes against rogue database edits", C_ORANGE)
    ]
    for i, (m_val, m_title, m_sub, m_col) in enumerate(metrics):
        x = Inches(0.8 + i * 3.95)
        mc = add_card(slide3, x, Inches(5.2), Inches(3.8), Inches(1.7), bg_color=RGBColor(15, 23, 42), border_color=RGBColor(30, 58, 95))
        tf_m = mc.text_frame
        tf_m.margin_left = Inches(0.2)
        tf_m.margin_top = Inches(0.15)
        p1 = tf_m.paragraphs[0]
        p1.text = m_val
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = m_col
        p2 = tf_m.add_paragraph()
        p2.text = m_title
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE
        p3 = tf_m.add_paragraph()
        p3.text = m_sub
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_MUTED

    # -------------------------------------------------------------
    # SLIDE 4: TECHNICAL ARCHITECTURE & SECURITY LAYERS
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "TECHNICAL ARCHITECTURE & SECURITY ENGINE",
               category="Under the Hood: Zero-Gas Consensus & Multi-Campus Security",
               subtext="A four-layer defense architecture guaranteeing data integrity, privacy, and speed.")

    layers = [
        ("Layer 1: Trust Layer",
         "Proof-of-Work Blockchain Ledger",
         "Custom PoW blockchain engine (SHA-256, Difficulty 4). Dual-persisted across SQLite block table and instance/blockchain.json. Zero gas fees for internal university consortium nodes.",
         C_CYAN),
        ("Layer 2: Cryptographic Layer",
         "Dynamic Salting & Zero Silent Passes",
         "PBKDF2-HMAC-SHA256 with 100,000 iterations and per-document dynamic cryptographic salting. Full blockchain ledger traversal re-derives digests from scratch on every verification.",
         C_GREEN),
        ("Layer 3: AI Document Ingestion",
         "Google Gemini 2.0 Flash AI",
         "Multimodal Vision AI extracts structured JSON from certificates (roll number, title, date, CGPA). Atomic all-or-nothing batch CSV validator halts invalid batches with 0 orphaned blocks.",
         C_PURPLE),
        ("Layer 4: Verification & UI",
         "Permanent Crest QR & Stamped PDF",
         "High-contrast QR with Error Correction Level H and embedded campus crest. Verification-grade Stamped PDF generator + interactive 3D WebGL ledger inspection.",
         C_ORANGE)
    ]

    for i, (l_tag, l_head, l_desc, l_col) in enumerate(layers):
        x = Inches(0.8 + i * 2.95)
        card = add_card(slide4, x, Inches(1.7), Inches(2.8), Inches(3.6), bg_color=C_CARD, border_color=C_CARD_BORDER)
        tf = card.text_frame
        tf.margin_left = Inches(0.18)
        tf.margin_top = Inches(0.18)
        tf.margin_right = Inches(0.18)

        p1 = tf.paragraphs[0]
        p1.text = l_tag
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = l_col

        p2 = tf.add_paragraph()
        p2.text = l_head
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        p3 = tf.add_paragraph()
        p3.text = l_desc
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_MUTED

    # Tech Stack Bar
    tech_banner = add_card(slide4, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.4), bg_color=RGBColor(15, 23, 42), border_color=C_CARD_BORDER)
    tf_tb = tech_banner.text_frame
    tf_tb.margin_left = Inches(0.25)
    tf_tb.margin_top = Inches(0.15)
    p_t1 = tf_tb.paragraphs[0]
    p_t1.text = "🛠️ PRODUCTION TECH STACK"
    p_t1.font.size = Pt(11)
    p_t1.font.bold = True
    p_t1.font.color.rgb = C_CYAN

    p_t2 = tf_tb.add_paragraph()
    p_t2.text = "• Backend: Python 3.14 + Flask 3.0 + SQLAlchemy + Flask-Login (Strict Role-Based Access Control)\n• Intelligence & Security: Google Gemini 2.0 Flash AI · PBKDF2-HMAC-SHA256 · Pillow + QR Error Correction H\n• Frontend: Responsive HTML5/CSS3 + Cyber Matrix WebGL 3D + Three.js Ledger Viewer + Desktop Native Mode\n• DevOps & Scale: Docker containerized + SQLite WAL mode + Polygon L2 rollup anchoring roadmap"
    p_t2.font.size = Pt(9.5)
    p_t2.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 5: LIVE PROTOTYPE WORKFLOW (THE 3 PORTALS)
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "LIVE PROTOTYPE: THREE ZERO-FRICTION PORTALS",
               category="End-to-End User Experience",
               subtext="Seamless user journey across University Registrar, Student, and Corporate Verifier.")

    portals = [
        ("🏛️ PORTAL 1: REGISTRAR / ISSUER",
         "AI-Assisted Credential Issuance",
         [
             "• Gemini 2.0 Flash AI auto-parses certificate PDFs into form fields in 2 seconds.",
             "• 1-Click Blockchain Anchoring: mines block, dynamic salts, and updates ledger.",
             "• Bulk CSV Issuance with atomic all-or-nothing pre-validation.",
             "• Automated email notification sent to student via SMTP gateway."
         ],
         C_CYAN),
        ("🎓 PORTAL 2: STUDENT WALLET",
         "Self-Sovereign Digital Ownership",
         [
             "• Permanent self-sovereign digital vault for all earned degrees and marksheets.",
             "• 3D Interactive Card Flip: inspect underlying block index and cryptographic hash.",
             "• Permanent Institutional QR Code embedded with official campus crest.",
             "• 1-Click Verification-Grade Stamped PDF Certificate download."
         ],
         C_GREEN),
        ("🔍 PORTAL 3: PUBLIC VERIFIER",
         "Instant 2-Second Verification",
         [
             "• Zero login barrier: employers simply scan QR or enter hash at /verify.",
             "• Real-time cryptographic audit trail: verified against ledger in under 2 seconds.",
             "• Anti-Tamper Engine: any altered database record trips instant RED alert.",
             "• Corporate Batch Verify portal for background checking multiple candidates."
         ],
         C_ORANGE)
    ]

    for i, (p_title, p_sub, p_bullets, p_col) in enumerate(portals):
        x = Inches(0.8 + i * 3.95)
        card = add_card(slide5, x, Inches(1.7), Inches(3.8), Inches(4.3), bg_color=C_CARD, border_color=C_CARD_BORDER)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)

        p1 = tf.paragraphs[0]
        p1.text = p_title
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = p_col

        p2 = tf.add_paragraph()
        p2.text = p_sub
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        for b in p_bullets:
            pb = tf.add_paragraph()
            pb.text = b
            pb.font.size = Pt(9.5)
            pb.font.color.rgb = C_MUTED

    # Bottom Prototype Status Banner
    p_banner = add_card(slide5, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.7), bg_color=RGBColor(15, 23, 42), border_color=C_GREEN)
    tf_pb = p_banner.text_frame
    tf_pb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_res = tf_pb.paragraphs[0]
    p_res.alignment = PP_ALIGN.CENTER
    p_res.text = "✅ 100% WORKING PROTOTYPE: Passed 11/11 Automated E2E Integration Suites | Dockerized | Tested at http://127.0.0.1:5000"
    p_res.font.size = Pt(11)
    p_res.font.bold = True
    p_res.font.color.rgb = C_GREEN

    # -------------------------------------------------------------
    # SLIDE 6: BUSINESS MODEL & MONETIZATION
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "BUSINESS MODEL & MONETIZATION STRATEGY",
               category="Commercial Viability & Unit Economics",
               subtext="A highly profitable, scalable B2B SaaS and pay-per-verify API revenue model.")

    streams = [
        ("🏢 B2B Institutional SaaS",
         "Annual University Subscription",
         "• Tiered annual subscription: ₹25 to ₹50 per student / year.\n• University benefits: Automated Gemini AI registrar ingestion, white-labeled student wallets, official branded QR crests, and administrative analytics.\n• Projected: 50 universities × 10,000 students = ₹2.0 Crores ARR.",
         C_CYAN),
        ("⚡ Corporate Pay-Per-Verify API",
         "Recruiter & BGV Screening API",
         "• Pay-per-call API for background verification giants (HireRight, First Advantage, AuthBridge, TCS, Infosys).\n• ₹15 to ₹30 per instant automated check (replacing ₹1,000+ manual checks).\n• Projected: 200,000 corporate checks/year = ₹40 Lakhs high-margin revenue.",
         C_GREEN),
        ("🏛️ Enterprise & Consortia",
         "State Board White-Label Deployments",
         "• Custom white-label deployment for State Higher Education Councils and Technical University Consortia.\n• Includes on-premise high-availability consortium clustering and custom ERP/LMS integrations.",
         C_PURPLE)
    ]

    for i, (st_tag, st_head, st_desc, st_col) in enumerate(streams):
        x = Inches(0.8 + i * 3.95)
        card = add_card(slide6, x, Inches(1.7), Inches(3.8), Inches(3.1), bg_color=C_CARD, border_color=C_CARD_BORDER)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)

        p1 = tf.paragraphs[0]
        p1.text = st_tag
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = st_col

        p2 = tf.add_paragraph()
        p2.text = st_head
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        p3 = tf.add_paragraph()
        p3.text = st_desc
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_MUTED

    # Unit Economics Card at Bottom
    econ_card = add_card(slide6, Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.9), bg_color=RGBColor(15, 23, 42), border_color=C_CARD_BORDER)
    tf_ec = econ_card.text_frame
    tf_ec.margin_left = Inches(0.25)
    tf_ec.margin_top = Inches(0.15)
    p_ec1 = tf_ec.paragraphs[0]
    p_ec1.text = "📊 UNBEATABLE UNIT ECONOMICS & PROFIT MARGINS"
    p_ec1.font.size = Pt(12)
    p_ec1.font.bold = True
    p_ec1.font.color.rgb = C_ORANGE

    p_ec2 = tf_ec.add_paragraph()
    p_ec2.text = "• Zero Gas Fees: Internal Proof-of-Work blockchain runs locally with ₹0.00 transaction fees.\n• L2 Polygon Anchoring: Periodic batch Merkle-tree anchoring onto Polygon costs only ~₹0.08 per credential.\n• Gross Profit Margin: >98% gross margin on corporate API verification requests.\n• Customer Acquisition Strategy: Direct campus onboarding starting with Jaypee Group, expanding to Tier-1 engineering colleges."
    p_ec2.font.size = Pt(10)
    p_ec2.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 7: MARKET OPPORTUNITY & COMPETITIVE ADVANTAGE
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "MARKET OPPORTUNITY & COMPETITIVE ADVANTAGE",
               category="Market Sizing & Competitive Matrix",
               subtext="Capturing India's ₹4,200 Cr credential verification market with clear technical superiority.")

    # Left: Market Sizing Cards
    market_cards = [
        ("TAM (Total Addressable Market)", "₹4,200 Crores", "Indian Background Screening, Document Fraud & Academic Credentialing Market.", C_CYAN),
        ("SAM (Serviceable Addressable Market)", "₹1,100 Crores", "40 Million+ students across 1,100+ Universities and 45,000+ Colleges.", C_GREEN),
        ("SOM (Serviceable Obtainable Market)", "₹45 Crores", "Targeting 300+ Private & Technical Universities in Phase 1 Rollout.", C_ORANGE)
    ]
    for i, (m_tag, m_val, m_desc, m_col) in enumerate(market_cards):
        y = Inches(1.7 + i * 1.6)
        card = add_card(slide7, Inches(0.8), y, Inches(4.3), Inches(1.45), bg_color=C_CARD, border_color=C_CARD_BORDER)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.12)
        p1 = tf.paragraphs[0]
        p1.text = m_tag
        p1.font.size = Pt(9.5)
        p1.font.bold = True
        p1.font.color.rgb = m_col

        p2 = tf.add_paragraph()
        p2.text = m_val
        p2.font.size = Pt(18)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        p3 = tf.add_paragraph()
        p3.text = m_desc
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = C_MUTED

    # Right: Competitive Matrix Table
    x_tbl = Inches(5.3)
    y_tbl = Inches(1.7)
    w_tbl = Inches(7.2)
    h_tbl = Inches(4.65)

    rows, cols = 6, 5
    table_shape = slide7.shapes.add_table(rows, cols, x_tbl, y_tbl, w_tbl, h_tbl)
    table = table_shape.table

    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(1.35)
    table.columns[2].width = Inches(1.25)
    table.columns[3].width = Inches(1.2)
    table.columns[4].width = Inches(1.2)

    headers_tbl = ["Feature / Metric", "Satyameva", "DigiLocker", "Traditional", "Public Web3"]
    matrix_data = [
        ["Verification Speed", "< 2 Seconds", "2 - 5 Minutes", "2 - 4 Weeks", "15 - 45 Sec"],
        ["Cost per Check", "Free / ₹15 API", "Free (Govt)", "₹500 - ₹2,000", "$2 - $15 Gas"],
        ["Employer Friction", "Zero (No Login)", "High (OTP Auth)", "High (Emails)", "High (Wallet)"],
        ["Internal DB Tamper Detect", "Cryptographic", "None (Trusts DB)", "None", "Cryptographic"],
        ["AI Registrar Automation", "Gemini 2.0 Flash", "None", "None (Manual)", "None"]
    ]

    for c, h in enumerate(headers_tbl):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(15, 23, 42)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = C_CYAN if c == 1 else C_WHITE
        p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT

    for r, row in enumerate(matrix_data):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_CARD if r % 2 == 0 else RGBColor(15, 23, 42)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.name = "Arial"
            if c == 1:
                p.font.bold = True
                p.font.color.rgb = C_GREEN
            elif c == 0:
                p.font.bold = True
                p.font.color.rgb = C_WHITE
            else:
                p.font.color.rgb = C_MUTED
            p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT

    # -------------------------------------------------------------
    # SLIDE 8: CAMPUS PILOT, ROADMAP & CONCLUSION
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "CAMPUS PILOT TRACTION & NATIONAL ROADMAP",
               category="Execution Roadmap & Vision",
               subtext="From a Jaypee multi-campus pilot to a nationwide decentralized trust protocol.")

    phases = [
        ("📍 PHASE 1: CAMPUS PILOT (Q4 2026)",
         "Jaypee Group Pilot Deployment",
         "• Live pilot across JUIT Waknaghat, JIIT Noida, and JUET Guna.\n• 15,000+ active students & 5,000+ annual graduating degrees.\n• Direct integration with campus registrar & training/placement cell.\n• Feedback iteration on UI and AI extraction models.",
         C_CYAN),
        ("🌐 PHASE 2: STATE EXPANSION (2027)",
         "DigiLocker & NAD Integration",
         "• REST API integration with DigiLocker and National Academic Depository (NAD).\n• Onboarding 25+ private and state technical universities.\n• Corporate partnerships with top 5 Indian background verification agencies.\n• Multi-tenant cloud clustering.",
         C_GREEN),
        ("🚀 PHASE 3: UNIVERSAL TRUST (2028+)",
         "National Scale & Beyond Academia",
         "• Scaling across 40M+ Indian higher-education students under NEP 2020.\n• Expanding to medical practitioner licenses, bar council credentials, and professional certifications.\n• Universal decentralized digital identity infrastructure.",
         C_PURPLE)
    ]

    for i, (ph_tag, ph_head, ph_desc, ph_col) in enumerate(phases):
        x = Inches(0.8 + i * 3.95)
        card = add_card(slide8, x, Inches(1.7), Inches(3.8), Inches(3.2), bg_color=C_CARD, border_color=C_CARD_BORDER)
        tf = card.text_frame
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_right = Inches(0.2)

        p1 = tf.paragraphs[0]
        p1.text = ph_tag
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = ph_col

        p2 = tf.add_paragraph()
        p2.text = ph_head
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_WHITE

        p3 = tf.add_paragraph()
        p3.text = ph_desc
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = C_MUTED

    # Grand Conclusion Banner
    close_banner = add_card(slide8, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.8), bg_color=RGBColor(15, 23, 42), border_color=C_CYAN)
    tf_cb = close_banner.text_frame
    tf_cb.margin_left = Inches(0.3)
    tf_cb.margin_top = Inches(0.2)

    p_c1 = tf_cb.paragraphs[0]
    p_c1.text = "🎯 TRANSFORMING ACADEMIC TRUST INTO MATHEMATICAL CERTAINTY"
    p_c1.font.size = Pt(15)
    p_c1.font.bold = True
    p_c1.font.color.rgb = C_CYAN

    p_c2 = tf_cb.add_paragraph()
    p_c2.text = "Satyameva Aegis eliminates weeks of manual friction, protects institutional credibility, and empowers students with lifetime self-sovereign ownership. We are ready to deploy."
    p_c2.font.size = Pt(11)
    p_c2.font.color.rgb = C_WHITE

    p_c3 = tf_cb.add_paragraph()
    p_c3.text = "TEAM: CYPHER-X   |   FOUNDER: Praneeth Sai   |   RIDEHACK'26 — IIC JUIT   |   GitHub: github.com/praneethsai050-cmd/satyameva"
    p_c3.font.size = Pt(9.5)
    p_c3.font.bold = True
    p_c3.font.color.rgb = C_MUTED

    # Save presentation
    output_path = "RIDEHACK_2026_Satyameva_Aegis.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    create_deck()
